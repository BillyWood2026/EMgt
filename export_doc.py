import docx

doc_path = r"e:\300软件开发\04战略管理体系\战略管理需求规格书（整合版）.docx"
doc = docx.Document(doc_path)

output = []
for i, para in enumerate(doc.paragraphs):
    style = para.style.name if para.style else ""
    text = para.text.strip()
    if text:
        prefix = ""
        if style == "Heading 1":
            prefix = "\n" + "=" * 80 + "\n"
        elif style == "Heading 2":
            prefix = "\n" + "-" * 60 + "\n"
        elif style == "Heading 3":
            prefix = "\n  "
        output.append(f"{prefix}[{style}] {text}")

# Also extract tables
output.append("\n\n" + "=" * 80)
output.append("TABLES")
output.append("=" * 80)
for t_idx, table in enumerate(doc.tables):
    output.append(f"\n--- Table {t_idx+1} ---")
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        output.append(" | ".join(cells))

with open(r"d:\EMgtv3\strategy_doc_full.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(output))

print(f"Exported {len(output)} lines to strategy_doc_full.txt")
