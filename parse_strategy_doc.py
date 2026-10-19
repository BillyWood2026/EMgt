import docx
import sys

doc_path = r"e:\300软件开发\04战略管理体系\战略管理需求规格书（整合版）.docx"
doc = docx.Document(doc_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print("=" * 80)

# Extract headings to understand structure
for i, para in enumerate(doc.paragraphs):
    style = para.style.name if para.style else ""
    text = para.text.strip()
    if not text:
        continue
    # Only print headings and non-empty paragraphs (limit to first 300 to understand structure)
    if i < 300:
        if "Heading" in style or "heading" in style or "标题" in style:
            print(f"[{i}] [{style}] {text}")
        elif text and len(text) < 100:
            # Print short paragraphs that might be section markers
            if any(text.startswith(p) for p in ["第", "一、", "二、", "三、", "四、", "五、", "六、", "七、", "八、", "九、", "十", "1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.", "10.", "11.", "12.", "13.", "14.", "15."]):
                print(f"[{i}] [{style}] {text}")
